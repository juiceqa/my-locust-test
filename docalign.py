from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

# Load your DOCX file
doc = Document("Garrett_Glick_Resume_2.docx")

# Map for readability of alignment
alignment_map = {
    0: "LEFT",
    1: "CENTER",
    2: "RIGHT",
    3: "JUSTIFY",
    None: "DEFAULT"
}

# Generate alignment report
print("Alignment Report:")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        alignment = para.alignment
        align_str = alignment_map.get(alignment, "UNKNOWN")
        print(f"{i}: [{align_str}] {text}")

# Define months to detect date lines
months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

# Correct alignment logic
for para in doc.paragraphs:
    text = para.text.strip()
    if not text:
        continue
    # Lines with tabs and months → right-aligned
    if '\t' in text and any(month in text for month in months):
        para.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    # Lines with tabs but no months → left-aligned (companies, job titles)
    elif '\t' in text:
        para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

# Save updated DOCX
doc.save("Garrett_Glick_Resume_Corrected.docx")
print("\nUpdated document saved as 'Garrett_Glick_Resume_Corrected.docx'")
